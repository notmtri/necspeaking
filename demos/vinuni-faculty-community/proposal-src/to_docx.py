"""Converts the proposal HTML (scratchpad/proposal/proposal.html) into an editable Word file."""
import json, re, sys
from pathlib import Path
from bs4 import BeautifulSoup, Comment, NavigableString, Tag
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor, Emu

P = Path(sys.argv[1]); OUT = sys.argv[2]
SHOTS = P / 'shots-print'; ASSETS = Path('/home/user/necspeaking/demos/vinuni-faculty-community/src/assets')
NAVY, NAVY_D, RED, INK, INK2, INK3 = '1D3C73', '0A1834', 'C72027', '1B1F28', '464E5E', '737B8C'
LINE, ZEBRA = 'DDE1E8', 'F7F8FA'
BOX = {'navy': ('EEF2FA', NAVY), 'red': ('FCECEF', RED), 'gold': ('FBF4E4', 'B88A2E'), 'green': ('E7F4EC', '1E7A4C')}
TAG = {'pilot': ('1E7A4C', 'PILOT'), 'later': ('464E5E', 'LATER'), 'chg': ('0E7A6E', None)}
SERIF, SANS = 'Georgia', 'Calibri'
TEXT_W = Cm(17.4)
rgb = lambda h: RGBColor.from_string(h)

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.left_margin = sec.right_margin = Cm(1.8); sec.top_margin = Cm(1.9); sec.bottom_margin = Cm(1.8)

# ---------- styles ----------
st = doc.styles
def font(style, name, size, color=INK, bold=None):
    f = style.font; f.name = name; f.size = Pt(size); f.color.rgb = rgb(color)
    if bold is not None: f.bold = bold
    rpr = style.element.get_or_add_rPr(); rf = rpr.find(qn('w:rFonts'))
    if rf is None: rf = OxmlElement('w:rFonts'); rpr.insert(0, rf)
    for a in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'): rf.set(qn(a), name)
font(st['Normal'], SANS, 10.5); st['Normal'].paragraph_format.space_after = Pt(5); st['Normal'].paragraph_format.line_spacing = 1.15
font(st['Heading 1'], SERIF, 22, NAVY_D, False); st['Heading 1'].paragraph_format.space_before = Pt(6); st['Heading 1'].paragraph_format.space_after = Pt(10)
font(st['Heading 2'], SERIF, 14, NAVY_D, False); st['Heading 2'].paragraph_format.space_before = Pt(14); st['Heading 2'].paragraph_format.space_after = Pt(5)
for s in ('Heading 1', 'Heading 2'): st[s].paragraph_format.keep_with_next = True
font(st['Caption'], SANS, 8.5, INK3, False); st['Caption'].font.italic = False; st['Caption'].paragraph_format.space_after = Pt(10)
for s in ('List Bullet', 'List Number'): font(st[s], SANS, 10.5); st[s].paragraph_format.space_after = Pt(3)
from docx.enum.style import WD_STYLE_TYPE
for s, size, color, indent in (('TOC 1', 10.5, INK, 0), ('TOC 2', 9.5, INK2, 0.6)):
    try: style = st[s]
    except KeyError: style = st.add_style(s, WD_STYLE_TYPE.PARAGRAPH); style.base_style = st['Normal']
    font(style, SANS, size, color); style.paragraph_format.left_indent = Cm(indent); style.paragraph_format.space_after = Pt(2)

# ---------- xml helpers ----------
def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr(); sh = OxmlElement('w:shd')
    sh.set(qn('w:val'), 'clear'); sh.set(qn('w:color'), 'auto'); sh.set(qn('w:fill'), fill); tcPr.append(sh)
def cell_borders(cell, **edges):
    tcPr = cell._tc.get_or_add_tcPr(); b = OxmlElement('w:tcBorders')
    for edge in ('top', 'left', 'bottom', 'right'):
        e = OxmlElement(f'w:{edge}'); spec = edges.get(edge)
        if spec: e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), str(spec[0])); e.set(qn('w:color'), spec[1])
        else: e.set(qn('w:val'), 'nil')
        b.append(e)
    tcPr.append(b)
def cell_margins(cell, top=80, bottom=80, left=120, right=120):
    tcPr = cell._tc.get_or_add_tcPr(); m = OxmlElement('w:tcMar')
    for k, v in (('top', top), ('start', left), ('bottom', bottom), ('end', right)):
        e = OxmlElement(f'w:{k}'); e.set(qn('w:w'), str(v)); e.set(qn('w:type'), 'dxa'); m.append(e)
    tcPr.append(m)
def no_table_borders(t):
    tblPr = t._tbl.tblPr; b = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        e = OxmlElement(f'w:{edge}'); e.set(qn('w:val'), 'nil'); b.append(e)
    tblPr.append(b)
def widths(t, ws):
    t.autofit = False
    tblPr = t._tbl.tblPr
    tw = tblPr.find(qn('w:tblW'))
    if tw is None: tw = OxmlElement('w:tblW'); tblPr.append(tw)
    tw.set(qn('w:w'), str(int(sum(ws) / 635))); tw.set(qn('w:type'), 'dxa')
    grid = t._tbl.tblGrid
    for gc, w in zip(grid.findall(qn('w:gridCol')), ws): gc.set(qn('w:w'), str(int(w / 635)))
    for row in t.rows:
        for c, w in zip(row.cells, ws): c.width = w
def keep_row(row):
    trPr = row._tr.get_or_add_trPr(); e = OxmlElement('w:cantSplit'); trPr.append(e)
def repeat_header(row):
    trPr = row._tr.get_or_add_trPr(); e = OxmlElement('w:tblHeader'); trPr.append(e)
def para_border_left(p, color):
    pPr = p._p.get_or_add_pPr(); b = OxmlElement('w:pBdr'); l = OxmlElement('w:left')
    l.set(qn('w:val'), 'single'); l.set(qn('w:sz'), '24'); l.set(qn('w:space'), '8'); l.set(qn('w:color'), color); b.append(l); pPr.append(b)
def field(par, instr, placeholder=''):
    r = par.add_run(); f1 = OxmlElement('w:fldChar'); f1.set(qn('w:fldCharType'), 'begin'); r._r.append(f1)
    r = par.add_run(); it = OxmlElement('w:instrText'); it.set(qn('xml:space'), 'preserve'); it.text = instr; r._r.append(it)
    r = par.add_run(); f2 = OxmlElement('w:fldChar'); f2.set(qn('w:fldCharType'), 'separate'); r._r.append(f2)
    if placeholder: par.add_run(placeholder)
    r = par.add_run(); f3 = OxmlElement('w:fldChar'); f3.set(qn('w:fldCharType'), 'end'); r._r.append(f3)

# ---------- containers ----------
class Box:
    """Wraps a Document or table cell; reuses a cell's empty first paragraph."""
    def __init__(self, target, width): self.t, self.width, self.fresh = target, width, hasattr(target, '_tc')
    def par(self, style=None):
        if self.fresh:
            self.fresh = False; p = self.t.paragraphs[0]
            if style: p.style = doc.styles[style]
            return p
        return self.t.add_paragraph(style=style)
    def table(self, rows, cols):
        self.fresh = False
        t = self.t.add_table(rows=rows, cols=cols)
        return t

# ---------- inline ----------
def inline(par, node, size=None, color=None, bold=False, italic=False):
    for ch in node.children:
        if isinstance(ch, NavigableString):
            txt = re.sub(r'\s+', ' ', str(ch))
            if not txt: continue
            r = par.add_run(txt); r.bold = bold or None; r.italic = italic or None
            if size: r.font.size = Pt(size)
            if color: r.font.color.rgb = rgb(color)
            continue
        if not isinstance(ch, Tag): continue
        cls = ch.get('class', [])
        if ch.name == 'br': par.add_run().add_break(); continue
        if ch.name in ('b', 'strong'): inline(par, ch, size, color or NAVY_D, True, italic); continue
        if ch.name in ('i', 'em'): inline(par, ch, size, color, bold, True); continue
        if ch.name == 'span' and 'tag' in cls:
            kind = next((k for k in TAG if k in cls), None)
            c, label = TAG.get(kind, (INK2, None))
            r = par.add_run('  ' + (label or ch.get_text(' ', strip=True).upper())); r.bold = True; r.font.size = Pt(7.5); r.font.color.rgb = rgb(c)
            continue
        if ch.name == 'span' and 'n' in cls:
            r = par.add_run(ch.get_text(strip=True) + '  '); r.font.color.rgb = rgb(RED); continue
        if ch.name == 'span' and ('muted' in cls): inline(par, ch, size, INK3, bold, italic); continue
        inline(par, ch, size, color, bold, italic)

def text_par(box, el, style=None, size=None, color=None, align=None, italic=False):
    p = box.par(style); inline(p, el, size, color, italic=italic)
    if align: p.alignment = align
    return p

# ---------- blocks ----------
pending_break = [False]
def heading(box, el, level):
    p = box.par(f'Heading {level}')
    if pending_break[0] and level == 1: p.paragraph_format.page_break_before = True; pending_break[0] = False
    inline(p, el)
    for r in p.runs:
        if r.font.color.rgb is None or str(r.font.color.rgb) != RED: r.font.color.rgb = rgb(NAVY_D)
        r.bold = False
    return p

def figure(box, el, width):
    img = el.find('img'); src = img['src'].replace('{{SHOTS}}', str(SHOTS))
    frac = next((int(c[1:]) / 100 for c in el.get('class', []) if re.fullmatch(r'w\d+', c)), 1.0)
    p = box.par(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER if frac < 1 else None
    p.paragraph_format.keep_with_next = True; p.paragraph_format.space_after = Pt(3)
    p.add_run().add_picture(src, width=int(width * frac))
    cap = el.find('figcaption')
    if cap:
        c = box.par('Caption'); inline(c, cap, color=INK3)
        if frac < 1: c.alignment = WD_ALIGN_PARAGRAPH.CENTER

def grid(box, el, ratios):
    kids = [k for k in el.children if isinstance(k, Tag)]
    t = box.table(1, len(kids)); no_table_borders(t)
    gap = Cm(0.5); avail = box.width - gap * (len(kids) - 1)
    ws = [int(avail * r / sum(ratios)) for r in ratios]
    widths(t, ws)
    for i, (k, w) in enumerate(zip(kids, ws)):
        c = t.rows[0].cells[i]; cell_margins(c, 0, 0, 0 if i == 0 else 140, 0 if i == len(kids) - 1 else 140)
        cb = Box(c, w - (Cm(0.25) if 0 < i < len(kids) - 1 else Cm(0.25)))
        if k.name == 'figure': figure(cb, k, cb.width)
        else: blocks(cb, k)
    box.par().paragraph_format.space_after = Pt(2)

def html_table(box, el):
    rows = el.find_all('tr'); ncol = max(len(r.find_all(['th', 'td'])) for r in rows)
    t = box.table(len(rows), ncol); no_table_borders(t)
    head = rows[0].find_all('th')
    fixed = []
    for c in (head or rows[0].find_all('td')):
        m = re.search(r'width:\s*(\d+)%', c.get('style', ''))
        fixed.append(int(m.group(1)) / 100 if m else ('num' in c.get('class', []) and 0.06) or None)
    free = 1 - sum(f for f in fixed if f); nfree = sum(1 for f in fixed if not f) or 1
    ws = [int(box.width * (f or free / nfree)) for f in fixed]
    widths(t, ws)
    for i, tr in enumerate(rows):
        keep_row(t.rows[i])
        cells = tr.find_all(['th', 'td'])
        if cells and cells[0].name == 'th': repeat_header(t.rows[i])
        for j, td in enumerate(cells):
            c = t.rows[i].cells[j]; cell_margins(c, 70, 70, 100, 100)
            th = td.name == 'th'
            if th: shade(c, NAVY)
            elif i % 2 == 0: shade(c, ZEBRA)
            cell_borders(c, bottom=(4, LINE))
            p = c.paragraphs[0]; p.paragraph_format.space_after = Pt(0)
            if 'num' in td.get('class', []): p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            inline(p, td, size=8.5 if th else 9, color='FFFFFF' if th else None, bold=th)
    box.par().paragraph_format.space_after = Pt(2)

def callout(box, el, kind):
    fill, edge = BOX[kind]
    t = box.table(1, 1); no_table_borders(t); widths(t, [box.width])
    c = t.rows[0].cells[0]; shade(c, fill); cell_borders(c, left=(36, edge)); cell_margins(c, 120, 120, 200, 160)
    blocks(Box(c, box.width - Cm(0.7)), el)
    box.par().paragraph_format.space_after = Pt(2)

def kpis(box, el):
    ks = el.select('.kpi'); t = box.table(1, len(ks)); no_table_borders(t)
    w = int((box.width - Cm(0.3) * (len(ks) - 1)) / len(ks)); widths(t, [w] * len(ks))
    for i, k in enumerate(ks):
        c = t.rows[0].cells[i]; shade(c, 'F3F5F9'); cell_margins(c, 120, 120, 140, 140)
        p = c.paragraphs[0]; r = p.add_run(k.select_one('.v').get_text(strip=True)); r.font.name = SERIF; r.font.size = Pt(22); r.font.color.rgb = rgb(NAVY_D); p.paragraph_format.space_after = Pt(2)
        q = c.add_paragraph(); inline(q, k.select_one('.l'), size=8.5, color=INK2); q.paragraph_format.space_after = Pt(0)
    box.par().paragraph_format.space_after = Pt(2)

def timeline(box, el):
    rows = el.select('.tl'); t = box.table(len(rows), 2); no_table_borders(t); widths(t, [Cm(3.2), box.width - Cm(3.2)])
    for i, tl in enumerate(rows):
        keep_row(t.rows[i])
        a, b = t.rows[i].cells
        for c in (a, b): cell_borders(c, bottom=(4, LINE)); cell_margins(c, 90, 90, 60, 60)
        p = a.paragraphs[0]; r = p.add_run(tl.select_one('.when').get_text(strip=True)); r.bold = True; r.font.color.rgb = rgb(RED); r.font.size = Pt(9.5)
        what = tl.select_one('.what'); title = what.find('b')
        q = b.paragraphs[0]; r = q.add_run(title.get_text(strip=True)); r.bold = True; r.font.color.rgb = rgb(NAVY_D); q.paragraph_format.space_after = Pt(1)
        title.extract(); s = b.add_paragraph(); inline(s, what, size=9.5); s.paragraph_format.space_after = Pt(0)
    box.par().paragraph_format.space_after = Pt(2)

def lists(box, el):
    style = 'List Number' if el.name == 'ol' else 'List Bullet'
    small = 'small' in el.get('class', [])
    first_par = [None]; last_pars = []
    for li in el.find_all('li', recursive=False):
        p = box.par(style); inline(p, li, size=9.5 if small else None)
        if first_par[0] is None: first_par[0] = p
        last_pars.append(p)
    if el.name == 'ol' and first_par[0] is not None:
        restart_numbering(first_par[0], last_pars)

def restart_numbering(_first, pars):
    numbering = doc.part.numbering_part.numbering_definitions._numbering
    style_num = doc.styles['List Number'].element.pPr.numPr.numId.val
    abstract = numbering.num_having_numId(style_num).abstractNumId.val
    num = numbering.add_num(abstract)
    ov = OxmlElement('w:lvlOverride'); ov.set(qn('w:ilvl'), '0'); so = OxmlElement('w:startOverride'); so.set(qn('w:val'), '1'); ov.append(so); num.append(ov)
    for p in pars:
        pPr = p._p.get_or_add_pPr(); numPr = pPr.get_or_add_numPr()
        numPr.get_or_add_ilvl().val = 0; numPr.get_or_add_numId().val = num.numId

def blocks(box, el):
    for ch in el.children:
        if not isinstance(ch, Tag):
            if isinstance(ch, NavigableString) and ch.strip(): text_par(box, BeautifulSoup(f'<span>{ch}</span>', 'html.parser').span)
            continue
        cls = ch.get('class', [])
        if ch.name == 'section':
            if 'newpage' in cls or 'flow' not in cls: pending_break[0] = True
            blocks(box, ch); continue
        if ch.name == 'h2': heading(box, ch, 1); continue
        if ch.name == 'h3': heading(box, ch, 2); continue
        if ch.name == 'h4':
            p = box.par(); r = p.add_run(ch.get_text(strip=True).upper()); r.bold = True; r.font.size = Pt(8.5); r.font.color.rgb = rgb(RED)
            p.paragraph_format.keep_with_next = True; p.paragraph_format.space_before = Pt(6); p.paragraph_format.space_after = Pt(2); continue
        if ch.name == 'p':
            if 'lede' in cls: text_par(box, ch, size=11.5, color=INK2)
            elif 'note' in cls or 'small' in cls or 'gate-note' in cls: text_par(box, ch, size=8.5, color=INK3)
            elif 'why' in cls: p = text_par(box, ch, size=9.5, color=INK2)
            else: text_par(box, ch)
            continue
        if ch.name in ('ul', 'ol'): lists(box, ch); continue
        if ch.name == 'table': html_table(box, ch); continue
        if ch.name == 'figure': figure(box, ch, box.width); continue
        if ch.name == 'div':
            if 'toc' in cls: toc(box); continue
            if 'box' in cls: callout(box, ch, next((k for k in BOX if k in cls), 'navy')); continue
            if 'kpis' in cls: kpis(box, ch); continue
            if 'timeline' in cls: timeline(box, ch); continue
            if 'figs' in cls:
                n = len([k for k in ch.children if isinstance(k, Tag)])
                grid(box, ch, [1.35, 1] if 'two-r' in cls else [1] * n); continue
            if 'side' in cls: grid(box, ch, [40, 60] if 'l' in cls else [56, 44]); continue
            if 'persona' in cls:
                p = box.par(); inline(p, ch.find_all('div', recursive=False)[-1]); p.paragraph_format.space_after = Pt(8)
                p.paragraph_format.keep_together = True; continue
            if 'why' in cls: blocks(box, ch); continue
            if 'sign' in cls:
                parts = ch.find_all('div', recursive=False); t = box.table(1, len(parts)); no_table_borders(t); widths(t, [int(box.width / len(parts))] * len(parts))
                for i, d in enumerate(parts):
                    c = t.rows[0].cells[i]; cell_borders(c, top=(6, INK3)); cell_margins(c, 60, 60, 0, 300)
                    c.paragraphs[0].add_run(d.get_text(strip=True)).font.size = Pt(9)
                continue
            blocks(box, ch); continue
        if ch.name == 'span': text_par(box, ch); continue

TOC_PAGES = json.loads((Path(OUT).with_suffix('.toc.json')).read_text()) if Path(OUT).with_suffix('.toc.json').exists() else {}
def fldchar(par, kind):
    r = par.add_run(); f = OxmlElement('w:fldChar'); f.set(qn('w:fldCharType'), kind); r._r.append(f)
def toc(box):
    heads = [(1 if h.name == 'h2' else 2, h.get_text(' ', strip=True)) for h in soup.select('h2.sec[id], h3[id]')]
    heads = [(lvl, re.sub(r'\s+(PILOT|LATER|Pilot|Later|Refinement \d)(\s+(Pilot|Later|Refinement \d))*$', '', t)) for lvl, t in heads]
    for i, (lvl, txt) in enumerate(heads):
        p = box.par(f'TOC {lvl}')
        p.paragraph_format.tab_stops.add_tab_stop(box.width, alignment=2, leader=1)
        p.paragraph_format.space_after = Pt(2 if lvl == 2 else 3)
        if i == 0:
            fldchar(p, 'begin'); r = p.add_run(); it = OxmlElement('w:instrText'); it.set(qn('xml:space'), 'preserve'); it.text = 'TOC \\o "1-2" \\h \\z \\u'; r._r.append(it); fldchar(p, 'separate')
        num, _, title = txt.partition(' ')
        a = p.add_run(num + '   '); a.font.color.rgb = rgb(RED); a.bold = lvl == 1
        b = p.add_run(title + '\t' + str(TOC_PAGES.get(num, ''))); b.bold = lvl == 1
        if i == len(heads) - 1: fldchar(p, 'end')
    n = box.par(); n.paragraph_format.space_before = Pt(6)
    z = n.add_run('After editing, right-click the list above and choose “Update Field” → “Update entire table” to refresh the page numbers.'); z.font.size = Pt(8.5); z.font.color.rgb = rgb(INK3)

# ---------- cover ----------
def cover():
    body = Box(doc, TEXT_W)
    t = doc.add_table(1, 2); no_table_borders(t); widths(t, [Cm(10), TEXT_W - Cm(10)]); body.fresh = False
    l, r = t.rows[0].cells
    l.paragraphs[0].add_run().add_picture(str(ASSETS / 'logo.png'), width=Cm(3.6))
    q = r.paragraphs[0]; q.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    x = q.add_run('PROPOSAL FOR A\nONE-SEMESTER PILOT'.replace('\n', ' ')); x.bold = True; x.font.size = Pt(8.5); x.font.color.rgb = rgb(RED)
    sp = doc.add_paragraph(); sp.paragraph_format.space_after = Pt(18)
    h = doc.add_paragraph(); r1 = h.add_run('International Faculty Community Platform'); r1.font.name = SERIF; r1.font.size = Pt(32); r1.font.color.rgb = rgb(NAVY_D)
    h.paragraph_format.space_after = Pt(8); h.paragraph_format.line_spacing = 1.0
    s = doc.add_paragraph(); r2 = s.add_run('Helping international faculty and their families settle in at VinUniversity, starting before they arrive.'); r2.font.size = Pt(13); r2.font.color.rgb = rgb(INK2)
    s.paragraph_format.space_after = Pt(16)
    doc.add_paragraph().add_run().add_picture(str(P / 'cover-lawn.jpg'), width=TEXT_W)
    meta = [('PREPARED FOR', [('School Council', True), ('VinUniversity', False)]), ('PREPARED BY', [('Nguyen Hoang Minh Tri', True), ('College of Engineering and Computer Science', False)]),
            ('BUILDS ON', [('International Faculty Community Platform proposal by Marc Kramer, Senior Lecturer (September 2026)', False)]), ('DATE', [('October 2026', True), ('Version 1.0, for discussion', False)])]
    t = doc.add_table(2, 2); no_table_borders(t); widths(t, [int(TEXT_W / 2)] * 2)
    for i, (lab, lines) in enumerate(meta):
        c = t.rows[i // 2].cells[i % 2]; cell_margins(c, 160, 100, 0, 200)
        p = c.paragraphs[0]; x = p.add_run(lab); x.bold = True; x.font.size = Pt(8); x.font.color.rgb = rgb(RED); p.paragraph_format.space_after = Pt(2)
        for txt, b in lines:
            q = c.add_paragraph(); y = q.add_run(txt); y.bold = b; y.font.size = Pt(10.5); y.font.color.rgb = rgb(NAVY_D if b else INK); q.paragraph_format.space_after = Pt(0)
    n = doc.add_paragraph(); n.paragraph_format.space_before = Pt(22)
    z = n.add_run('Accompanies the interactive demo. All people and providers shown in screenshots are fictional sample data.'); z.font.size = Pt(8.5); z.font.color.rgb = rgb(INK3)

# ---------- footer ----------
def footer():
    for s in doc.sections:
        f = s.footer.paragraphs[0]; f.text = ''
        from docx.enum.text import WD_TAB_ALIGNMENT
        for pos in (Emu(4680 * 635), Emu(9360 * 635)): f.paragraph_format.tab_stops.add_tab_stop(pos, alignment=WD_TAB_ALIGNMENT.CLEAR)
        f.paragraph_format.tab_stops.add_tab_stop(TEXT_W, alignment=WD_TAB_ALIGNMENT.RIGHT)
        r = f.add_run('International Faculty Community Platform · Pilot proposal\t'); r.font.size = Pt(8); r.font.color.rgb = rgb(INK3)
        field(f, 'PAGE', '1')
        for x in f.runs: x.font.size = Pt(8); x.font.color.rgb = rgb(INK3)
    sec.different_first_page_header_footer = True

html = (P / 'proposal.html').read_text()
soup = BeautifulSoup(html, 'html.parser')
for c in soup.find_all(string=lambda t: isinstance(t, Comment)): c.extract()
cover()
pending_break[0] = True
blocks(Box(doc, TEXT_W), soup.body)
footer()
zoom = doc.settings.element.find(qn('w:zoom'))
if zoom is not None: zoom.set(qn('w:percent'), '100')
cp = doc.core_properties; cp.title = 'International Faculty Community Platform: Proposal for a One-Semester Pilot'; cp.author = 'Nguyen Hoang Minh Tri'
TBL = ['tblStyle', 'tblpPr', 'tblOverlap', 'bidiVisual', 'tblStyleRowBandSize', 'tblStyleColBandSize', 'tblW', 'jc', 'tblCellSpacing', 'tblInd', 'tblBorders', 'shd', 'tblLayout', 'tblCellMar', 'tblLook', 'tblCaption', 'tblDescription']
TC = ['cnfStyle', 'tcW', 'gridSpan', 'hMerge', 'vMerge', 'tcBorders', 'shd', 'noWrap', 'tcMar', 'textDirection', 'tcFitText', 'vAlign', 'hideMark']
TR = ['cnfStyle', 'divId', 'gridBefore', 'gridAfter', 'wBefore', 'wAfter', 'cantSplit', 'trHeight', 'tblHeader', 'tblCellSpacing', 'jc', 'hidden']
def reorder(parent, order):
    kids = list(parent); key = lambda e: order.index(e.tag.split('}')[1]) if e.tag.split('}')[1] in order else len(order)
    for e in kids: parent.remove(e)
    for e in sorted(kids, key=key): parent.append(e)
body = doc.element.body
for e in body.iter(qn('w:tblPr')): reorder(e, TBL)
for e in body.iter(qn('w:tcPr')): reorder(e, TC)
for e in body.iter(qn('w:trPr')): reorder(e, TR)
doc.save(OUT); print('saved', OUT)
